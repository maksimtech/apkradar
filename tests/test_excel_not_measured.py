"""A row that was not analysed keeps its name and shows no counts.

Measured on 2026-10-09 with a real registry of eleven F-Droid APKs plus one row
naming an app and a package whose file was not there. `batch-excel` wrote that
row as

    None | None | APK | None | (blank) | N/A | 0 | 0 | 0 | None | None

The name and the package the registry gave it were gone — the scan result of a
file that could not be opened has neither, and the report took them from the
result rather than from the row — so the one line a reader most needs to
identify was the one line that could not be identified. And "0 trackers, 0
permissions, 0 transfers" for a file that was never opened is a measurement of
zero where nothing was measured, in the columns that get summed and charted.
The score cell was already left blank for exactly that reason; the counts and
the format were not. The SKIPPED row (package name, no file) had the same zeros.
"""

from __future__ import annotations

import openpyxl
import pytest
from typer.testing import CliRunner

from apkradar.cli import app
from apkradar.excel import write_results
from apkradar.scanner import ScanResult
from tests.android_files import manifest, write_apk

runner = CliRunner()


def _registry(tmp_path, apk_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["App Name", "Package Name", "APK Path"])
    ws.append(["Mastodon", "org.joinmastodon.android", apk_path])
    ws.append(["Coin Master", "com.moonactive.coinmaster", None])
    ws.append(["Missing App", "com.example.missing", str(tmp_path / "not-there.apk")])
    src = tmp_path / "registry.xlsx"
    wb.save(src)
    return src


@pytest.fixture
def report(tmp_path):
    apk = write_apk(tmp_path / "mastodon.apk", manifest("org.joinmastodon.android"))
    src = _registry(tmp_path, apk)
    out = tmp_path / "report.xlsx"
    outcome = runner.invoke(app, ["batch-excel", str(src), "--output", str(out)])
    assert outcome.exit_code == 1, outcome.output     # one file was missing, and that is said
    return openpyxl.load_workbook(out).worksheets[0]


@pytest.fixture
def augmented(tmp_path):
    apk = write_apk(tmp_path / "mastodon.apk", manifest("org.joinmastodon.android"))
    src = _registry(tmp_path, apk)
    outcome = runner.invoke(app, ["batch-excel", str(src), "--augment"])
    assert outcome.exit_code == 1, outcome.output
    return openpyxl.load_workbook(src).worksheets[0]


def _row(ws, key, column=0):
    rows = list(ws.iter_rows(min_row=2, values_only=True))
    for row in rows:
        if row[column] == key:
            return row
    raise AssertionError(f"no row with {key!r} in column {column}: {[r[column] for r in rows]}")


def test_the_missing_file_row_keeps_the_name_and_package_the_registry_gave_it(report):
    row = _row(report, "com.example.missing", column=1)
    assert row[0] == "Missing App"
    assert row[5] == "N/A"


def test_the_missing_file_row_has_no_format(report):
    """It was never opened, so it has no format — "APK" was the dataclass default."""
    assert _row(report, "com.example.missing", column=1)[2] in (None, "")


@pytest.mark.parametrize("package", ["com.example.missing", "com.moonactive.coinmaster"])
def test_nothing_measured_means_no_counts(report, package):
    row = _row(report, package, column=1)
    assert row[4] in (None, "")                   # Score, as before
    assert row[6:9] == (None, None, None) or row[6:9] == ("", "", "")   # Trackers, Permissions, Extra-EU


def test_a_measured_row_still_has_its_zeros(report):
    # The report names the app as the APK labels it, which here is its package.
    row = _row(report, "org.joinmastodon.android", column=1)
    assert row[4] == 100
    assert row[6:9] == (0, 0, 0)


def test_the_augmented_sheet_leaves_the_same_cells_blank(augmented):
    header = [c.value for c in augmented[1]]
    trackers = header.index("Trackers")
    for name in ("Missing App", "Coin Master"):
        row = _row(augmented, name)
        assert row[trackers:trackers + 3] in ((None, None, None), ("", "", ""))
    assert _row(augmented, "Mastodon")[trackers:trackers + 3] == (0, 0, 0)


def test_write_results_alone_does_the_same(tmp_path):
    out = tmp_path / "out.xlsx"
    write_results(
        [ScanResult(apk_path="gone.apk", package_name="com.example.gone", error="File not found: gone.apk")],
        str(out),
    )
    ws = openpyxl.load_workbook(out).worksheets[0]
    assert ws["B2"].value == "com.example.gone"
    assert ws["C2"].value in (None, "")
    assert (ws["G2"].value, ws["H2"].value, ws["I2"].value) in ((None, None, None), ("", "", ""))
