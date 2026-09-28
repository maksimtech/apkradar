"""
APKRadar — l'impronta del pacchetto va mostrata intera.

Il referto stampava `SHA256:  af0db3dc4a74b14b...`: sedici caratteri su
sessantaquattro. Un'impronta serve a una cosa sola, farsi verificare, e
sedici caratteri non permettono di verificare niente — non si confronta con
l'hash pubblicato dallo store, non si cerca in un elenco, non si riporta in
un allegato. Chi vuole controllare deve ricalcolarla e non ha con cosa.

Due incoerenze mostravano il difetto già prima di cercarlo:

- la lettera al DPO (`templates/dpo_letter_it.txt`) usa `{{ sha256 }}` intero,
  quindi il terzo che riceve la contestazione ha l'impronta completa e
  l'analista che deve verificarla no;
- le citazioni di legge stampano i 64 caratteri dell'hash del testo applicato.
  Lo strumento era più rigoroso con la formulazione della norma che con la
  prova.

Il foglio Excel tronca allo stesso modo, ed è la consegna che resta in mano a
chi legge dopo: lì una colonna larga non costa niente.
"""
from __future__ import annotations

from unittest.mock import patch

from typer.testing import CliRunner

from apkradar.cli import app
from apkradar.scanner import ScanResult

runner = CliRunner()

# Un digest vero di 64 caratteri: quello del pacchetto base di 112 Where ARE U
# non serve, serve che la lunghezza sia quella giusta e che non si ripeta.
DIGEST = "af0db3dc4a74b14b9c2e7f1d05a836be41cc9df2e8b7a05614d3f2c8b9071e5a"


def _scanned() -> ScanResult:
    return ScanResult(
        apk_path="demo.apk",
        package_name="com.example.demo",
        app_name="Demo",
        version_name="1.0",
        version_code="1",
        sha256=DIGEST,
        apk_format="apk",
    )


def test_the_report_shows_the_whole_digest():
    with patch("apkradar.scanner.scan", return_value=_scanned()):
        result = runner.invoke(app, ["audit", "demo.apk"])

    assert result.exit_code == 0, result.output
    # Senza ritorni a capo in mezzo: un'impronta spezzata su due righe si
    # incolla male tanto quanto una troncata.
    assert DIGEST in result.output.replace("\n", "")


def test_the_report_does_not_offer_a_truncated_digest():
    """Il troncamento con i puntini sembra un'impronta e non lo è."""
    with patch("apkradar.scanner.scan", return_value=_scanned()):
        result = runner.invoke(app, ["audit", "demo.apk"])

    assert f"{DIGEST[:16]}..." not in result.output


def test_the_spreadsheet_cell_holds_the_whole_digest(tmp_path):
    from openpyxl import load_workbook

    from apkradar.excel import write_results

    target = tmp_path / "referto.xlsx"
    write_results([_scanned()], str(target))

    book = load_workbook(target)
    cells = {
        str(cell.value)
        for sheet in book.worksheets
        for row in sheet.iter_rows()
        for cell in row
        if cell.value
    }
    assert DIGEST in cells, "il foglio non porta l'impronta intera"
